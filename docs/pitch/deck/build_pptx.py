"""Build the EDITABLE Zill-and-Shams-pitch.pptx: native text boxes, rectangles and rules; pictures only for the
photo layers (photo-only backgrounds, one per photo slide); speaker notes on every slide.
From frontend/:  node ../docs/pitch/deck/extract.mjs   (writes deck/pptx/layout.json + bg-NN.png from index.html)
Then:            backend/.venv/Scripts/python.exe docs/pitch/deck/build_pptx.py
Fonts: Georgia (Fraunces) and Calibri (Inter Tight) so it renders on any machine; Arabic in Arial."""
import io, glob, os
from PIL import Image
from pptx import Presentation
from pptx.util import Emu
here = os.path.dirname(os.path.abspath(__file__))
qa = os.path.join(here, '..', 'qa')
NOTES = [
"Zill & Shams (ظل وشمس, shade and sun). An AI agent that plans walks and rides around the sun. Team Abhamzay, Team 13. Live at zill.itshamza.me.",
"Maps optimise time. Hot cities: route choice decides sun exposure. Dark winters: too little daylight (Oslo about 6 h at the December solstice). Sun and shade move every minute, so static lists go stale.",
"Sourced figures only. 72% of 12,346 Abu Dhabi adults vitamin D deficient (BMC Research Notes 2019); 70,000+ riders (Dubai RTA); 18.72 M visitors 2024 (Dubai DET via WAM); UAE midday break 12:30-15:00, 15 Jun-15 Sep (MoHRE 2025).",
"One engine, two modes. Shade mode = least direct sun (includes AC links like Metro to Dubai Mall). Sun mode = most direct sun for daylight and well-being. The preference flips.",
"Open map data, 3D model, sun and shadows every 10 minutes, best route, agent. Dubai: 8,602 buildings, 7,659 street segments, 744 trees. Oslo is the second city, built by the same pipeline.",
"Shade mode proof: Burj Park to Business Bay Metro, summer 15:30. 6.2 min less direct sun, 40% less, 2.2 min longer. Model estimate; awnings and wind not modelled. Source: data/scenarios/hero.json.",
"Sun mode proof, real routes from the documented search (hero.json 'sun'). Dubai 15 Apr 10:30: sun 1,101 s to 1,625 s (+524 s, +48%) for +269 s. Oslo 21 Jun 12:00: 613 s to 1,288 s (+675 s, +110%) for +131 s. Clear-sky geometry, model estimates.",
"Oslo: 13,841 buildings in the shadow ring, 10,082 edges, 20,208 trees, 53 POIs (meta.json). Sun ~6.6 deg at midwinter noon, shadows ~8.6x height, tested. Winter weather in the app is a labelled stand-in. Images are shade-bake masks.",
"Well-being: WHO UV band, indicative heat stress, daylight window (sun high, UV >= 3, cloud < 70%), break prompts. Not medical advice. In Oslo winter, sun mode is about daylight and mood, not vitamin D (UV 0.1 in the solstice scenario).",
"Agent plans, checks (verifier; numbers come from engine tokens), and adapts. Pluggable LLM chain OpenCode Go then Voke. Without a live key, sessions are labelled recorded replays.",
"Accuracy and freshness (docs/ACCURACY.md): sun within 0.5 deg of pvlib on 20 samples; shadows 10 min / 2 m; weather hourly; map re-bake on demand, weekly job planned. About 79% of Dubai heights are estimates. No field validation yet, so no accuracy percentage claimed. 470 automated tests pass (375 backend, 95 frontend; data/test_results.json).",
"Security (docs/SECURITY.md): no accounts or stored user data; location per request only; per-IP rate limits, size caps, CORS allow-list, CSP/HSTS, log scrubbing, key server-side. pip-audit clean; npm audit 2 findings in Next's PostCSS. PDPL/GDPR alignment is an aim, not certification. Pen-test comes before B2B launch.",
"Who pays. Target customers (not partners): Talabat, Deliveroo, Careem. Expanded scope: city planners (B2G planning analytics: where to plant trees or add shade), schools, elderly care, events, real estate. Pricing hypotheses (not validated): hotels ~$99 per property per month; fleets ~$1 per rider per month; outdoor employers ~$49 per site per month; city licence ~$15,000 per year; corporate well-being ~$12 per employee per year; consumer premium ~$2.99 per month, no ads.",
"Freemium, no ads. Reason: ads would put location data into an ad pipeline; no ads keeps the privacy story clean. Free app, optional premium (saved routes, day plans), B2B/B2G subscriptions as main revenue, pay-per-route API.",
"Costs (docs/COSTS.md). Server $6.93/month (Contabo VPS, 4 vCPU, 8 GB); per 1,000 users about $0.002 of server; fixed cost about $17/month. OpenCode Go $10/month flat: ~3,630 conversations on glm-5.3-flash, ~56 on kimi-k3 (page request estimates, 8.7 requests per conversation measured). Bake timings on a dev laptop, not the VM.",
"Users. Market: 70,000+ riders, 18.72 M visitors (sourced; context, not a forecast). Capacity measured with scripts/loadtest.py: local machine peaks ~26 req/s at 20 users (p95 596 ms), 40 users p95 2.6 s; live VM at 20 users 16.9 req/s, p95 1.6 s, 0 errors (run capped at 20 users, did not saturate). VM: 4 vCPU, 8 GB.",
"Roadmap: Dubai and Oslo now; UAE and Gulf next with city planners as first B2G pilot; northern cities winter daylight pilot. No funding ask. Team Abhamzay: Mohammed Hamza Ansari, Abhay Preethan. Try it at zill.itshamza.me.",
]
import json
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Pt
from pptx.oxml.ns import qn
from lxml import etree
LAY = json.load(open(os.path.join(here, 'pptx', 'layout.json'), encoding='utf8'))
assert len(LAY) == len(NOTES), (len(LAY), len(NOTES))
PX = 6350
def E(v): return Emu(int(round(v * PX)))
FONTS = {'Fraunces': 'Georgia', 'Inter Tight': 'Calibri', 'Amiri': 'Arial'}
FSCALE = {'Georgia': 0.93, 'Calibri': 0.97, 'Arial': 0.95}
def set_alpha(fill_fmt, a):
    sf = fill_fmt._xPr.find(qn('a:solidFill'))
    clr = sf[0]
    el = etree.SubElement(clr, qn('a:alpha')); el.set('val', str(int(a * 100000)))
def rect(slide, x, y, w, h, c, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, E(x), E(y), E(max(w, 0.5)), E(max(h, 0.5)))
    sp.fill.solid(); sp.fill.fore_color.rgb = RGBColor(c['r'], c['g'], c['b'])
    if c['a'] < 0.999: set_alpha(sp.fill, c['a'])
    sp.line.fill.background(); sp.shadow.inherit = False
    return sp
def textbox(slide, t):
    runs_all = [r for p in t['paras'] for r in p]
    sz0 = next((r['size'] for r in runs_all if 'size' in r), 24)
    tb = slide.shapes.add_textbox(E(t['x']), E(t['y']), E(t['w'] * (1.0 if t['single'] else 1.03)), E(t['h']))
    tf = tb.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = not t['single']; tf.vertical_anchor = MSO_ANCHOR.TOP
    para = tf.paragraphs[0]
    para.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[t['align']]
    para.line_spacing = Pt(t['lh'] * 0.5 * 0.97)
    for r in runs_all:
        if r.get('br'):
            para.add_line_break(); continue
        run = para.add_run(); run.text = ('   ' + r['t']) if r['t'].startswith('“') and para.runs and para.runs[-1].font.name == 'Arial' else r['t']
        fam = FONTS.get(r['font'], 'Calibri')
        f = run.font; f.name = fam; f.size = Pt(max(6, r['size'] * 0.5 * FSCALE.get(fam, 1) * (0.9 if fam == 'Georgia' and r['size'] >= 60 else 1)))
        f.bold = r['weight'] >= 600; f.italic = r.get('italic', False)
        f.color.rgb = RGBColor(r['color']['r'], r['color']['g'], r['color']['b'])
        rPr = run._r.get_or_add_rPr()
        if r['ls']: rPr.set('spc', str(int(r['ls'] * 0.5 * 100)))
        if r['color']['a'] < 0.999:
            sf = rPr.find(qn('a:solidFill'))
            el = etree.SubElement(sf[0], qn('a:alpha')); el.set('val', str(int(r['color']['a'] * 100000)))
        if fam == 'Arial':
            rPr.set('lang', 'ar-AE'); cs = etree.SubElement(rPr, qn('a:cs')); cs.set('typeface', 'Arial')
prs = Presentation(); prs.slide_width = Emu(12191695); prs.slide_height = Emu(6858000)
blank = prs.slide_layouts[6]
for i, (lay, note) in enumerate(zip(LAY, NOTES), 1):
    s = prs.slides.add_slide(blank)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor(0x0E, 0x0C, 0x0A)
    bg = os.path.join(here, 'pptx', f'bg-{i:02d}.png')
    if lay['bake'] and os.path.exists(bg):
        buf = io.BytesIO(); Image.open(bg).convert('RGB').save(buf, 'JPEG', quality=90); buf.seek(0)
        s.shapes.add_picture(buf, 0, 0, prs.slide_width, prs.slide_height)
    for sh in lay['shapes']:
        if sh['kind'] == 'fill': rect(s, sh['x'], sh['y'], sh['w'], sh['h'], sh['color'])
        elif sh['kind'] == 'dot': rect(s, sh['x'], sh['y'], sh['w'], sh['h'], sh['color'], MSO_SHAPE.OVAL)
        else:
            bw, side = sh['bw'], sh['side']
            if side == 'Top': rect(s, sh['x'], sh['y'], sh['w'], bw, sh['color'])
            elif side == 'Bottom': rect(s, sh['x'], sh['y'] + sh['h'] - bw, sh['w'], bw, sh['color'])
            elif side == 'Left': rect(s, sh['x'], sh['y'], bw, sh['h'], sh['color'])
            else: rect(s, sh['x'] + sh['w'] - bw, sh['y'], bw, sh['h'], sh['color'])
    for t in lay['texts']: textbox(s, t)
    s.notes_slide.notes_text_frame.text = note
prs.save(os.path.join(here, '..', 'Zill-and-Shams-pitch.pptx'))
print('ok', len(LAY))
